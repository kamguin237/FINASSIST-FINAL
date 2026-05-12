from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "ANALYSE_DU_SYSTEME_FINASSIST_MAJ.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(text)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)


def build_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)

    styles["Title"].font.name = "Calibri"
    styles["Title"].font.size = Pt(21)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(31, 78, 121)

    styles["Heading 1"].font.name = "Calibri"
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)

    styles["Heading 2"].font.name = "Calibri"
    styles["Heading 2"].font.size = Pt(12.5)
    styles["Heading 2"].font.bold = True
    styles["Heading 2"].font.color.rgb = RGBColor(55, 86, 35)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("ANALYSE DU SYSTÈME - PROJET FINASSIST")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run(
        "Version mise à jour selon l'implémentation réelle observée dans le dépôt au 27 avril 2026"
    )
    run.italic = True

    note = doc.add_paragraph()
    note.paragraph_format.space_after = Pt(12)
    note.add_run("Note méthodologique : ").bold = True
    note.add_run(
        "cette version met à jour l'analyse textuelle du système à partir du code existant. "
        "Les schémas UML référencés doivent être redessinés ou ajustés si l'on souhaite une "
        "cohérence parfaite avec cette analyse révisée."
    )

    doc.add_heading("1. Introduction à l’analyse", level=1)
    add_body(
        doc,
        "La phase d'analyse consiste à étudier en détail les besoins fonctionnels, les règles "
        "métier et l'organisation du système afin de modéliser son comportement avant ou pendant "
        "le développement. Dans le cas de FINASSIST, l'implémentation actuelle montre un système "
        "centré sur la gestion des besoins internes, leur validation hiérarchique, la signature "
        "électronique, les notifications, le reporting et l'administration des accès.",
    )
    add_body(doc, "L'analyse UML pertinente pour l'application effectivement implémentée repose sur :")
    for item in [
        "le diagramme de cas d'utilisation ;",
        "le diagramme de séquence ;",
        "le diagramme d'activité ;",
        "le diagramme d'état-transition.",
    ]:
        add_bullet(doc, item)
    add_body(doc, "Ces diagrammes permettent de représenter respectivement :")
    for item in [
        "les interactions entre les acteurs et le système ;",
        "le déroulement temporel des échanges entre couches applicatives ;",
        "la logique des processus métier ;",
        "le cycle de vie des entités et des sessions.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("2. Diagrammes de Cas d’Utilisation", level=1)
    doc.add_heading("2.1 Objectif", level=2)
    add_body(
        doc,
        "Le diagramme de cas d'utilisation permet de structurer les besoins fonctionnels en "
        "identifiant précisément qui fait quoi dans l'application. Dans FINASSIST, il met en "
        "évidence une séparation nette entre création des besoins, validation, signature, "
        "administration et supervision.",
    )
    add_body(doc, "La version implémentée couvre notamment :")
    for item in [
        "la gestion des besoins internes, de leur création à leur clôture ;",
        "la validation hiérarchique par circuits paramétrables ;",
        "la signature électronique des documents et la gestion de signature personnelle ;",
        "l'administration des utilisateurs, rôles, permissions, catégories et workflows ;",
        "la consultation du dashboard, des rapports et des journaux d'activité ;",
        "la gestion des notifications et des préférences utilisateur.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("2.2 Acteurs identifiés", level=2)
    add_body(doc, "Les acteurs fonctionnels effectivement identifiés dans le système sont :")
    for item in [
        "Administrateur ;",
        "Direction ;",
        "Responsable ;",
        "Agent / Employé.",
    ]:
        add_bullet(doc, item)
    add_body(
        doc,
        "Le code montre également des rôles techniques secondaires via les permissions, mais ces "
        "quatre acteurs constituent les profils métier structurants visibles dans les parcours.",
    )

    doc.add_heading("2.3 Description des principaux cas d’utilisation", level=2)
    add_body(doc, "Gestion des besoins")
    for item in [
        "l'agent crée un besoin, ajoute des pièces jointes, l'enregistre et le soumet ;",
        "l'agent consulte le détail, le statut, l'historique et les documents associés ;",
        "le responsable ou la direction consulte les besoins en attente de leur niveau ;",
        "le valideur approuve, rejette ou transmet selon les règles du workflow ;",
        "le système applique les règles de délai, les notifications et les changements d'état.",
    ]:
        add_bullet(doc, item)

    add_body(doc, "Gestion de la signature électronique")
    for item in [
        "l'utilisateur gère sa signature personnelle ;",
        "la signature peut être manuscrite, typographique, importée ou capturée via QR code mobile ;",
        "le signataire appose sa signature sur le document du besoin ;",
        "le système enregistre l'empreinte, l'horodatage, le PDF signé et la vérification d'authenticité.",
    ]:
        add_bullet(doc, item)

    add_body(doc, "Administration du système")
    for item in [
        "l'administrateur gère les utilisateurs, rôles et permissions ;",
        "il configure les catégories et les circuits de workflow ;",
        "il consulte les logs et les états du système ;",
        "il supervise les modules sans être acteur de validation métier.",
    ]:
        add_bullet(doc, item)

    add_body(doc, "SCHEMA DU DIAGRAMME")
    add_body(
        doc,
        "Figure 1 à mettre à jour : diagramme de cas d'utilisation global de FINASSIST "
        "(gestion des besoins, validation, signature, notifications, reporting et administration).",
    )

    doc.add_heading("2.4 Importance du diagramme de cas d’utilisation", level=2)
    for item in [
        "il clarifie les responsabilités de chaque acteur ;",
        "il définit les limites du système réellement implémenté ;",
        "il sert de base aux diagrammes de séquence, d'activité et d'état ;",
        "il révèle le rôle central des permissions dans l'accès aux fonctionnalités.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("3. Diagrammes de Séquence", level=1)
    doc.add_heading("3.1 Objectif", level=2)
    add_body(
        doc,
        "Le diagramme de séquence décrit l'enchaînement chronologique des interactions entre "
        "l'utilisateur, le frontend Angular, l'API ASP.NET Core, les services applicatifs et la "
        "base SQL Server. Il met en lumière la traduction technique des règles métier.",
    )
    add_body(doc, "Dans FINASSIST, il permet notamment de détailler :")
    for item in [
        "la connexion et la sécurisation par JWT ;",
        "la création, l'enregistrement et la soumission d'un besoin ;",
        "la validation hiérarchique et la transmission entre étapes ;",
        "la signature électronique et la restitution du document signé ;",
        "la génération des notifications et le traitement des délais.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("3.2 Description des principaux diagrammes de séquence", level=2)
    add_body(doc, "Connexion (Login)")
    for item in [
        "l'utilisateur saisit son email et son mot de passe ;",
        "le frontend envoie une requête `POST /api/auth/login` ;",
        "le backend vérifie l'utilisateur, contrôle le mot de passe haché et génère le JWT ;",
        "le système journalise la connexion avec informations techniques et géographiques ;",
        "le frontend stocke le contexte utilisateur et ouvre l'accès aux modules autorisés.",
    ]:
        add_number(doc, item)
    add_body(doc, "SCHEMA DU DIAGRAMME")
    add_body(doc, "Figure 2 : diagramme de séquence de connexion à aligner avec le backend `AuthController` et `AuthService`.")

    add_body(doc, "Création et soumission d’un besoin")
    for item in [
        "l'agent remplit le formulaire et peut joindre un document ;",
        "le frontend appelle `POST /api/besoins` puis, si besoin, l'ajout de pièce jointe ;",
        "le besoin est créé en base avec son statut initial ;",
        "l'agent peut l'enregistrer, puis le soumettre au circuit de validation ;",
        "le système positionne l'étape courante, historise l'action et notifie les validateurs concernés.",
    ]:
        add_number(doc, item)
    add_body(doc, "SCHEMA DU DIAGRAMME")
    add_body(doc, "Figure 3 : diagramme de séquence de création / soumission d’un besoin.")

    add_body(doc, "Validation hiérarchique")
    for item in [
        "le valideur consulte les besoins à traiter ;",
        "il transmet une décision APPROUVE ou REJETE via l'API workflow ;",
        "le backend contrôle le rôle requis, interdit la double validation et exige un motif en cas de rejet ;",
        "le besoin change d'état, l'historique est mis à jour et des notifications sont envoyées ;",
        "si l'étape exige une signature, la progression dépend ensuite de la signature du document.",
    ]:
        add_number(doc, item)
    add_body(doc, "SCHEMA DU DIAGRAMME")
    add_body(doc, "Figure 4 : diagramme de séquence du workflow de validation.")

    add_body(doc, "Signature électronique")
    for item in [
        "le signataire récupère le document associé au besoin ;",
        "la signature personnelle est utilisée ou capturée pour produire une signature exploitable ;",
        "le frontend peut générer un PDF signé, puis l'API enregistre l'empreinte et la signature ;",
        "le backend met à jour le statut, historise l'action et permet la vérification de l'authenticité ;",
        "le document signé peut ensuite être téléchargé.",
    ]:
        add_number(doc, item)
    add_body(doc, "SCHEMA DU DIAGRAMME")
    add_body(doc, "Figure 5 : diagramme de séquence de la signature électronique.")

    add_body(doc, "Notifications et rappels")
    for item in [
        "la soumission, la transmission, le rejet ou la signature déclenchent une notification métier ;",
        "les notifications sont stockées en base puis diffusées selon les préférences utilisateur ;",
        "le système peut envoyer du push web via VAPID ;",
        "un worker traite périodiquement les délais pour générer rappels, emails et rejet automatique si nécessaire.",
    ]:
        add_number(doc, item)
    add_body(doc, "SCHEMA DU DIAGRAMME")
    add_body(doc, "Figure 6 : diagramme de séquence des notifications et du traitement automatique des délais.")

    doc.add_heading("4. Diagrammes d’Activité", level=1)
    doc.add_heading("4.1 Objectif", level=2)
    add_body(
        doc,
        "Le diagramme d'activité représente le flux opérationnel des processus métiers. "
        "Dans FINASSIST, il formalise surtout le cycle de vie d'un besoin et les décisions "
        "associées à son traitement.",
    )
    add_body(doc, "Il met en évidence :")
    for item in [
        "les étapes successives du processus ;",
        "les décisions conditionnelles ;",
        "les cas de rejet ou de transmission ;",
        "les conséquences d'une signature requise ;",
        "les effets des délais et rappels automatiques.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("4.2 Traitement d’un besoin", level=2)
    add_body(doc, "Le processus métier implémenté peut être résumé ainsi :")
    for item in [
        "création du besoin ;",
        "enregistrement éventuel en brouillon ou état intermédiaire ;",
        "soumission au circuit de validation ;",
        "prise en charge par le valideur du niveau courant ;",
        "approbation ou rejet ;",
        "signature à l'étape si elle est requise ;",
        "transmission manuelle à l'étape suivante ;",
        "clôture par signature finale ou rejet définitif.",
    ]:
        add_number(doc, item)
    add_body(
        doc,
        "Un flux complémentaire doit être représenté pour les dépassements de délai : rappel à 50 %, "
        "rappel à 80 %, email à l'expiration et rejet automatique après dépassement prolongé.",
    )
    add_body(doc, "SCHEMA DU DIAGRAMME")
    add_body(doc, "Figure 7 : diagramme d’activité du traitement d’un besoin.")

    doc.add_heading("5. Diagrammes d'État-Transition", level=1)
    doc.add_heading("5.1 Objectif", level=2)
    add_body(
        doc,
        "Le diagramme d'état-transition permet de modéliser le cycle de vie d'une entité ou d'une "
        "session. Dans FINASSIST, il est particulièrement utile pour la session utilisateur, le besoin "
        "métier, la session QR de signature et le document signé.",
    )

    doc.add_heading("5.2 Description des diagrammes d'état-transition", level=2)
    add_body(doc, "Session utilisateur")
    add_body(
        doc,
        "Le frontend implémente une logique de session avec authentification, surveillance "
        "d'inactivité, avertissement utilisateur et déconnexion automatique. Les états peuvent être "
        "modélisés comme suit : non connecté, authentification en cours, connecté, avertissement "
        "d'inactivité, session expirée.",
    )
    add_body(doc, "SCHEMA DU DIAGRAMME")
    add_body(doc, "Figure 8 : diagramme d’état-transition de session utilisateur.")

    add_body(doc, "Cycle de vie d’un besoin")
    add_body(
        doc,
        "Le code fait apparaître un automate métier plus fin que la version initiale du document. "
        "Les états observés ou dérivables sont : `BROUILLON`, `ENREGISTRE`, `SOUMISE`, "
        "`EN_ATTENTE_<ROLE>`, `APPROUVE_PAR_<ROLE>`, `SIGNE_PAR_<ROLE>`, `TRANSMIS`, "
        "`REJETE_PAR_<ROLE>` et état final de terminaison lorsque la dernière signature aboutit.",
    )
    add_body(
        doc,
        "Les transitions sont déclenchées par la soumission, la décision du valideur, la signature, "
        "la transmission manuelle ou le rejet automatique pour dépassement de délai.",
    )
    add_body(doc, "SCHEMA DU DIAGRAMME")
    add_body(doc, "Figure 9 : diagramme d’état-transition du besoin.")

    add_body(doc, "Session QR de signature")
    add_body(
        doc,
        "Une session QR suit également un cycle de vie spécifique : session créée, session ouverte "
        "sur mobile, signature soumise, session complétée ou session expirée.",
    )
    add_body(doc, "SCHEMA DU DIAGRAMME")
    add_body(doc, "Figure 10 : diagramme d’état-transition d’une session QR de signature.")

    doc.add_heading("5.3 Importance des diagrammes d'état-transition", level=2)
    for item in [
        "ils clarifient le cycle de vie réel des entités métier ;",
        "ils rendent explicites les règles de transition implémentées dans le backend ;",
        "ils servent de support pour les validations, la signature et les traitements automatiques ;",
        "ils renforcent la traçabilité et la compréhension des états dynamiques.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("6. Synthèse technique de l’analyse", level=1)
    add_body(doc, "Les éléments structurels observés dans l’implémentation peuvent être résumés comme suit :")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Élément"
    hdr[1].text = "Analyse issue du code"
    set_cell_shading(hdr[0], "1F4E79")
    set_cell_shading(hdr[1], "1F4E79")
    for cell in hdr:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.bold = True

    rows = [
        ("Frontend", "Application Angular modulaire avec guards, services, layout dynamique, notifications et gestion d’inactivité."),
        ("Backend", "API ASP.NET Core organisée en couches API, Application, Core et Infrastructure."),
        ("Sécurité", "JWT, hachage Argon2, permissions fines par rôles et affectations directes."),
        ("Persistance", "SQL Server avec Entity Framework Core, entités métier et migrations."),
        ("Workflow", "Circuits paramétrables, étapes ordonnées, rôles requis, signature requise et délais."),
        ("Notifications", "Notifications en base, push web VAPID, emails de rappel et filtres par préférences utilisateur."),
        ("Audit", "Logs d’activité, historique des besoins, journalisation de connexion et métadonnées techniques."),
    ]
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    doc.add_heading("7. Conclusion de l’analyse", level=1)
    add_body(
        doc,
        "L’analyse révisée montre que FINASSIST est un système de gestion des besoins internes, "
        "adossé à un workflow configurable, à une signature électronique multi-parcours et à un "
        "socle d’administration et d’audit relativement complet.",
    )
    add_body(doc, "Les diagrammes UML à conserver ou à mettre à jour doivent désormais refléter :")
    for item in [
        "la centralité du module Besoins ;",
        "la dynamique du workflow par rôles et permissions ;",
        "la signature électronique, y compris via QR mobile ;",
        "les notifications et la gestion automatique des délais ;",
        "les préférences utilisateur et la gestion de session côté interface.",
    ]:
        add_bullet(doc, item)
    add_body(
        doc,
        "Cette analyse constitue ainsi une base plus fidèle à l’implémentation réellement livrée "
        "et peut être utilisée pour mettre en cohérence la conception, la documentation et les "
        "schémas de mémoire.",
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    end.add_run("FIN DU DOCUMENT").bold = True

    doc.core_properties.title = "Analyse du système FINASSIST - mise à jour"
    doc.core_properties.subject = "Analyse UML alignée sur l'implémentation"
    doc.core_properties.author = "Codex"
    doc.core_properties.comments = (
        "Document révisé à partir du document initial et du dépôt FINASSIST."
    )
    doc.core_properties.created = datetime(2026, 4, 27, 12, 0, 0)
    return doc


def main() -> None:
    doc = build_document()
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
