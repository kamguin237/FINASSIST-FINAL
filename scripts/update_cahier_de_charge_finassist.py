from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "CAHIER_DE_CHARGE_FINASSIST_MAJ.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(text)


def build_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)

    styles["Title"].font.name = "Calibri"
    styles["Title"].font.size = Pt(22)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(31, 78, 121)

    styles["Heading 1"].font.name = "Calibri"
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)

    styles["Heading 2"].font.name = "Calibri"
    styles["Heading 2"].font.size = Pt(12.5)
    styles["Heading 2"].font.bold = True
    styles["Heading 2"].font.color.rgb = RGBColor(55, 86, 35)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("CAHIER DE CHARGE")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run(
        "Application FINASSIST\nVersion mise à jour selon le code implémenté au 27 avril 2026"
    )
    run.italic = True

    note = doc.add_paragraph()
    note.paragraph_format.space_after = Pt(14)
    note.add_run("Note de mise à jour : ").bold = True
    note.add_run(
        "la présente version reflète le périmètre fonctionnel et technique observé dans le dépôt "
        "FINASSIST. Les éléments de planning initial qui ne sont pas déductibles du code ont été "
        "résumés et signalés comme points à revalider."
    )

    doc.add_heading("INTRODUCTION", level=1)

    doc.add_heading("CONTEXTE", level=2)
    add_body(
        doc,
        "Dans un environnement financier où la traçabilité, la rapidité de traitement et le "
        "contrôle interne sont essentiels, FINASSIST vise à digitaliser la gestion des besoins "
        "internes et de leur circuit d'approbation.",
    )
    add_body(
        doc,
        "Le code actuellement implémenté montre que l'application s'est recentrée sur un noyau "
        "fonctionnel clair : création de besoins internes, validation hiérarchique paramétrable, "
        "signature électronique, notifications, reporting et administration des accès.",
    )
    add_body(doc, "Les limites adressées par la solution sont notamment :")
    for item in [
        "les retards liés aux traitements manuels ou partiellement numérisés ;",
        "le manque de visibilité sur l'état d'avancement de chaque besoin ;",
        "la difficulté à tracer les validations, transmissions, signatures et rejets ;",
        "les risques d'erreurs, d'oubli ou de perte d'information ;",
        "la nécessité d'un meilleur contrôle des habilitations et de l'audit.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("PRÉSENTATION GÉNÉRALE DU BESOIN", level=2)
    add_body(
        doc,
        "FINASSIST est une application web centralisée destinée à gérer les besoins internes d'une "
        "institution, depuis leur création jusqu'à leur validation finale et, lorsque requis, leur "
        "signature électronique.",
    )
    add_body(
        doc,
        "Le périmètre réellement implémenté dans le code ne fait pas apparaître un module autonome "
        "de réclamations. Le cœur du produit est la gestion des besoins internes, structurés par "
        "catégories et pilotés par des circuits de workflow configurables.",
    )
    add_body(doc, "La solution permet notamment :")
    for item in [
        "la saisie, l'enregistrement et la soumission des besoins ;",
        "la gestion des catégories et de leur rattachement à un circuit de validation ;",
        "la validation hiérarchique avec décisions, commentaires, rejets motivés et transmission manuelle ;",
        "la signature électronique de documents avec conservation et vérification ;",
        "la gestion des notifications in-app, push web et rappels par email ;",
        "la consultation d'indicateurs, rapports, exports et journaux d'activité ;",
        "l'administration fine des utilisateurs, rôles, permissions et préférences.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("PÉRIMÈTRE FONCTIONNEL IMPLÉMENTÉ", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Domaine"
    hdr[1].text = "Constat dans l'implémentation"
    set_cell_shading(hdr[0], "1F4E79")
    set_cell_shading(hdr[1], "1F4E79")
    for cell in hdr:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.bold = True

    rows = [
        ("Besoins", "Module principal implémenté avec brouillon, enregistrement, soumission, détail, historique et pièces jointes."),
        ("Workflow", "Circuits paramétrables par catégorie, étapes ordonnées, rôles requis, délais, validation, rejet et transmission."),
        ("Signature", "Signature personnelle utilisateur, signature de document, vérification d'authenticité, PDF signé et signature via QR mobile."),
        ("Notifications", "Notifications applicatives, push web via VAPID, rappels automatiques et emails de dépassement de délai."),
        ("Reporting", "Dashboard, statistiques, évolution par période, rapport filtré et export PDF/Excel."),
        ("Administration", "Gestion des utilisateurs, rôles, permissions, activation/désactivation et permissions directes."),
        ("Audit", "Logs d'activité, historique des besoins, journalisation des connexions et actions sensibles."),
        ("Paramètres", "Préférences utilisateur, langue, format de date, fuseau horaire, pagination, page d'accueil et paramètres de session."),
    ]
    for domaine, constat in rows:
        cells = table.add_row().cells
        cells[0].text = domaine
        cells[1].text = constat

    doc.add_heading("OBJECTIFS DU LOGICIEL", level=2)
    add_body(doc, "L'objectif principal de FINASSIST est de fiabiliser et d'automatiser le traitement des besoins internes en offrant :")
    for item in [
        "une création structurée et traçable des besoins ;",
        "un workflow dynamique basé sur des circuits paramétrables par catégorie ;",
        "une validation hiérarchique avec règles métier et contrôle des habilitations ;",
        "une signature électronique intégrée au traitement documentaire ;",
        "un suivi complet de l'historique, des notifications et des logs ;",
        "un tableau de bord opérationnel et des rapports exportables ;",
        "une administration centralisée des utilisateurs, rôles et permissions ;",
        "des rappels d'échéance et un traitement automatique des retards de validation.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("ACTEURS / UTILISATEURS", level=2)

    add_body(doc, "Les rôles effectivement présents dans le code sont les suivants :")

    add_body(doc, "Administrateur")
    for item in [
        "crée, modifie, désactive, réactive et supprime les utilisateurs ;",
        "gère les rôles, les permissions et les affectations directes ;",
        "gère les catégories et les circuits de workflow ;",
        "consulte les logs et les modules d'administration ;",
        "n'intervient pas comme valideur métier dans le workflow des besoins.",
    ]:
        add_bullet(doc, item)

    add_body(doc, "Direction")
    for item in [
        "consulte les besoins en attente de son niveau ;",
        "approuve ou rejette les besoins transmis à son étape ;",
        "signe les documents lorsque l'étape l'exige ;",
        "consulte le tableau de bord, le reporting et les notifications selon ses permissions.",
    ]:
        add_bullet(doc, item)

    add_body(doc, "Responsable")
    for item in [
        "consulte les besoins relevant de son niveau de validation ;",
        "prend une décision d'approbation ou de rejet ;",
        "signe le document à son étape si la signature est requise ;",
        "transmet manuellement le besoin à l'étape suivante.",
    ]:
        add_bullet(doc, item)

    add_body(doc, "Agent / Employé")
    for item in [
        "crée un besoin et le modifie tant qu'il n'est pas engagé irréversiblement ;",
        "ajoute, consulte et supprime des pièces jointes ;",
        "enregistre ou soumet son besoin ;",
        "consulte le statut, l'historique, les documents signés et les notifications ;",
        "gère sa signature personnelle, y compris via QR code mobile ;",
        "paramètre ses préférences d'application et de session.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("ANALYSE DES BESOINS", level=1)

    doc.add_heading("Besoins fonctionnels", level=2)

    add_body(doc, "1. Authentification et sécurité")
    for item in [
        "connexion par email et mot de passe ;",
        "émission et validation d'un jeton JWT ;",
        "déconnexion applicative ;",
        "protection des routes frontend et des endpoints backend ;",
        "contrôle d'accès par permissions métier.",
    ]:
        add_number(doc, item)

    add_body(doc, "2. Gestion des utilisateurs, rôles et permissions")
    for item in [
        "création, modification, activation, désactivation et suppression d'utilisateurs ;",
        "changement de rôle et changement de mot de passe ;",
        "création, modification et suppression de rôles ;",
        "consultation et mise à jour des permissions ;",
        "attribution de permissions à un rôle ou directement à un utilisateur ;",
        "consultation des permissions effectives.",
    ]:
        add_number(doc, item)

    add_body(doc, "3. Gestion des besoins")
    for item in [
        "création d'un besoin avec catégorie, niveau d'importance et pièces jointes ;",
        "mise à jour du besoin tant que son état le permet ;",
        "enregistrement intermédiaire et soumission au workflow ;",
        "consultation de la liste des besoins selon le contexte utilisateur ;",
        "consultation du détail, des documents et de l'historique ;",
        "suppression du besoin lorsqu'aucune contrainte métier ne l'interdit ;",
        "suivi des échéances de validation.",
    ]:
        add_number(doc, item)

    add_body(doc, "4. Workflow de validation")
    for item in [
        "création de circuits composés de plusieurs étapes ordonnées ;",
        "rattachement d'une catégorie à un circuit ;",
        "définition d'un rôle requis, d'une signature requise et d'un délai maximum par étape ;",
        "gestion des décisions APPROUVE et REJETE ;",
        "motif obligatoire en cas de rejet ;",
        "interdiction de double validation par le même utilisateur ;",
        "transmission manuelle à l'étape suivante ;",
        "rappels de délai, email de dépassement et rejet automatique en cas de retard prolongé ;",
        "gestion de statuts dynamiques de type EN_ATTENTE_ROLE, APPROUVE_PAR_ROLE, SIGNE_PAR_ROLE et REJETE_PAR_ROLE.",
    ]:
        add_number(doc, item)

    add_body(doc, "5. Signature électronique")
    for item in [
        "création, mise à jour et suppression d'une signature personnelle ;",
        "prise en charge des modes manuscrit, typographique, upload d'image et QR mobile ;",
        "signature électronique des documents d'un besoin ;",
        "génération et téléchargement d'un PDF signé ;",
        "vérification d'authenticité et d'intégrité de la signature.",
    ]:
        add_number(doc, item)

    add_body(doc, "6. Notifications et communication")
    for item in [
        "notifications lors de la soumission, transmission, signature, rejet et rappel ;",
        "consultation et marquage comme lu des notifications ;",
        "envoi manuel de notifications par les utilisateurs autorisés ;",
        "notifications push web via service worker et VAPID ;",
        "emails automatiques de rappel en cas de délai dépassé.",
    ]:
        add_number(doc, item)

    add_body(doc, "7. Reporting, tableau de bord et audit")
    for item in [
        "consultation du dashboard et des indicateurs synthétiques ;",
        "visualisation de l'évolution des besoins par période ;",
        "consultation des statistiques globales ;",
        "recherche de rapports filtrés par dates et statut ;",
        "export des rapports en PDF et Excel ;",
        "consultation des logs d'activité et des historiques de besoins.",
    ]:
        add_number(doc, item)

    add_body(doc, "8. Paramètres utilisateur")
    for item in [
        "gestion des préférences de notifications ;",
        "choix de langue, format de date et fuseau horaire ;",
        "configuration de la pagination, du tri par défaut et de la page d'accueil ;",
        "paramètres de session : déconnexion automatique, durée d'inactivité et avertissement.",
    ]:
        add_number(doc, item)

    doc.add_heading("Besoins non fonctionnels", level=2)
    for title, items in [
        (
            "Sécurité",
            [
                "authentification JWT ;",
                "mots de passe hachés avec Argon2 ;",
                "gestion des permissions fines ;",
                "journalisation des actions sensibles et des connexions.",
            ],
        ),
        (
            "Performance",
            [
                "architecture séparée frontend / backend / base de données ;",
                "chargement ciblé des modules côté interface ;",
                "exports et reporting adaptés à la volumétrie métier.",
            ],
        ),
        (
            "Fiabilité",
            [
                "migrations de base de données ;",
                "worker de surveillance des délais ;",
                "gestion des erreurs métier et techniques ;",
                "conservation des historiques, signatures et logs.",
            ],
        ),
        (
            "Maintenabilité",
            [
                "architecture en couches côté backend : API, Application, Core, Infrastructure ;",
                "frontend Angular structuré par modules, services, guards et composants partagés ;",
                "conteneurisation locale via Docker Compose.",
            ],
        ),
        (
            "Ergonomie",
            [
                "interface web responsive ;",
                "navigation par modules et permissions ;",
                "accès mobile au parcours de signature via QR code.",
            ],
        ),
    ]:
        add_body(doc, title)
        for item in items:
            add_bullet(doc, item)

    doc.add_heading("TECHNOLOGIES ET ARCHITECTURE", level=1)

    add_body(doc, "Architecture logicielle observée")
    for item in [
        "architecture 3 tiers : frontend web, API backend, base de données ;",
        "backend en ASP.NET Core Web API avec séparation en quatre projets : API, Application, Core et Infrastructure ;",
        "frontend Angular basé sur des composants standalone, guards, interceptors et services ;",
        "base de données SQL Server avec Entity Framework Core et migrations ;",
        "orchestration locale via Docker Compose (SQL Server, backend, frontend).",
    ]:
        add_bullet(doc, item)

    add_body(doc, "Technologies effectivement utilisées")
    for item in [
        "ASP.NET Core Web API, Entity Framework Core, JWT, QuestPDF, ClosedXML, WebPush ;",
        "Angular, RxJS, Chart.js, ngx-toastr, ng2-pdf-viewer, pdf-lib, qrcode ;",
        "SQL Server ;",
        "Docker et Docker Compose ;",
        "SMTP pour les emails de rappel ;",
        "service worker pour les notifications push web.",
    ]:
        add_bullet(doc, item)

    add_body(doc, "Point d'attention")
    add_bullet(
        doc,
        "une classe FirebaseNotificationService existe dans le code mais sous forme de stub ; l'envoi push réellement exploité repose sur Web Push avec VAPID.",
    )

    doc.add_heading("RESSOURCES", level=1)

    add_body(doc, "Ressources humaines mobilisables")
    for item in [
        "chef de projet / coordination ;",
        "développement backend .NET ;",
        "développement frontend Angular ;",
        "administration base de données SQL Server ;",
        "tests fonctionnels, sécurité et validation métier.",
    ]:
        add_bullet(doc, item)

    add_body(doc, "Ressources logicielles")
    for item in [
        "Visual Studio ou VS Code ;",
        "Git ;",
        "Postman ;",
        "Docker Desktop ;",
        "SQL Server ;",
        "outils de maquette si nécessaire.",
    ]:
        add_bullet(doc, item)

    add_body(doc, "Ressources matérielles")
    for item in [
        "poste de développement compatible Docker ;",
        "accès réseau interne ou internet selon le mode de déploiement ;",
        "serveur de déploiement ou infrastructure cloud ;",
        "mécanisme de sauvegarde et de sécurisation des certificats/secret keys.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("PLANNING", level=1)
    add_body(
        doc,
        "Le planning détaillé figurant dans la version initiale du cahier de charge était prévisionnel. "
        "Le code source permet de confirmer les modules livrés, mais ne permet pas de reconstituer "
        "de manière fiable les dates réelles d'exécution, les charges ou les écarts de planning.",
    )
    add_body(doc, "En conséquence, la mise à jour du document porte principalement sur :")
    for item in [
        "le périmètre fonctionnel effectivement implémenté ;",
        "la stack technique réellement utilisée ;",
        "les ajouts introduits pendant le développement, notamment QR mobile, préférences utilisateur, notifications push et gestion automatisée des délais.",
    ]:
        add_bullet(doc, item)
    add_body(
        doc,
        "Le planning projet doit être revalidé à partir des sources de suivi de projet "
        "(journal d'avancement, tickets, commits datés, comptes rendus ou planning de stage).",
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    annex = doc.add_paragraph()
    annex.alignment = WD_ALIGN_PARAGRAPH.CENTER
    annex.add_run("FIN DU DOCUMENT").bold = True

    doc.core_properties.title = "Cahier de charge FINASSIST - mise à jour"
    doc.core_properties.subject = "Version alignée sur l'implémentation réelle"
    doc.core_properties.author = "Codex"
    doc.core_properties.comments = (
        "Document révisé à partir du cahier initial et du code source disponible dans le dépôt FINASSIST."
    )
    doc.core_properties.created = datetime(2026, 4, 27, 12, 0, 0)

    return doc


def main() -> None:
    doc = build_document()
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
